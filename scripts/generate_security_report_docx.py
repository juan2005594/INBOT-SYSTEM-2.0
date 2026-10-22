import os
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def create_report():
    doc = Document()

    # Configurar márgenes
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # PORTADA
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("\n\n\nINFORME TÉCNICO DE AUDITORÍA Y PRUEBAS DE SEGURIDAD")
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(16, 44, 87)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("Evaluación de Controles Técnicos, API REST y Gestión de Acceso\nSistema de Inventario INBOT SYSTEM")
    run_sub.font.size = Pt(14)
    run_sub.font.color.rgb = RGBColor(53, 95, 142)

    doc.add_paragraph("\n" * 4)

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.add_run("UNIVERSIDAD DE CUNDINAMARCA\n").bold = True
    p_meta.add_run("Campo de Aprendizaje: MEDIT\nCADI: Seguridad en Aplicaciones\n\n")
    p_meta.add_run("Auditor / Desarrollador: Pablo Andrés Sánchez\n")
    p_meta.add_run("Fecha: Octubre 2026\n")

    doc.add_page_break()

    # TABLA DE CONTENIDO / INTRODUCCIÓN
    h1 = doc.add_heading("1. Introducción", level=1)
    doc.add_paragraph(
        "El presente informe técnico expone los resultados integrales de la auditoría de seguridad y "
        "las pruebas técnicas ejecutadas sobre la plataforma INBOT SYSTEM (API REST para gestión de inventarios, "
        "ventas y compras). El objetivo primordial consiste en validar el cumplimiento de los controles de "
        "autenticación, autorización, validación de entradas, gestión de sesiones, manejo seguro de archivos y "
        "exposición de servicios web, identificando vulnerabilidades potenciales y formulando un plan estructurado de mitigación."
    )

    # 2. OBJETIVOS
    doc.add_heading("2. Objetivos", level=1)
    doc.add_heading("2.1 Objetivo General", level=2)
    doc.add_paragraph(
        "Evaluar mediante pruebas técnicas y funcionales los principales controles de seguridad de la aplicación "
        "INBOT SYSTEM, documentando evidencias empíricas, resultados, hallazgos de riesgo y recomendaciones de mejora."
    )
    doc.add_heading("2.2 Objetivos Específicos", level=2)
    doc.add_paragraph("• Auditar dependencias de backend y frontend identificando versiones obsoletas o vulnerables.", style='List Bullet')
    doc.add_paragraph("• Inspeccionar la configuración de seguridad del framework y el despliegue (cookies, SSL, secretos).", style='List Bullet')
    doc.add_paragraph("• Evaluar los mecanismos de autenticación (Tokens DRF y sesiones) y control de acceso basado en roles (RBAC).", style='List Bullet')
    doc.add_paragraph("• Ejecutar pruebas positivas y negativas sobre endpoints REST mediante solicitudes automatizadas y Postman.", style='List Bullet')
    doc.add_paragraph("• Evaluar los controles en la recepción y procesamiento de imágenes y archivos multimedia.", style='List Bullet')

    # 3. DESCRIPCIÓN DE LA APLICACIÓN
    doc.add_heading("3. Descripción de la Aplicación", level=1)
    doc.add_paragraph(
        "INBOT SYSTEM es una solución integral orientada al control de inventarios, proveedores, facturas de compras "
        "y registro de ventas con soporte para interacción móvil y chatbot de asistencia. Su arquitectura técnica comprende:"
    )
    doc.add_paragraph("• Backend Principal: Python 3.13 / Django 5.0.14 y Django REST Framework (DRF 3.16.0).", style='List Bullet')
    doc.add_paragraph("• Base de Datos: Microsoft SQL Server (instancia ING_SANCHEZ, BD: inbotf_system) vía mssql-django y pyodbc (con soporte secundario SQLite).", style='List Bullet')
    doc.add_paragraph("• Capa Móvil / Chatbot: Módulo aplicacion_movil con validación perceptual de imágenes de producto (Pillow + pyzbar) y sesiones de chat.", style='List Bullet')
    doc.add_paragraph("• Frontend / Plantillas: Interfaz web renderizada con Django Templates, HTML5 y JavaScript Vanilla.", style='List Bullet')

    # 4. HERRAMIENTAS Y COMANDOS DE AUDITORÍA
    doc.add_heading("4. Herramientas y Comandos de Auditoría", level=1)
    doc.add_paragraph(
        "La metodología de auditoría se adaptó al ecosistema tecnológico de la solución (Python/Django y Node.js/React), "
        "ejecutando las herramientas equivalentes a los estándares de la guía académica:"
    )
    
    table_tools = doc.add_table(rows=1, cols=3)
    table_tools.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table_tools.rows[0].cells
    hdr[0].text = "Herramienta / Comando"
    hdr[1].text = "Ecosistema / Equivalente"
    hdr[2].text = "Propósito en la Auditoría"
    for c in hdr:
        set_cell_background(c, "102C57")
        c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        c.paragraphs[0].runs[0].font.bold = True

    tool_rows = [
        ("python manage.py check --deploy", "Django Deployment Security", "Auditoría de directivas críticas de seguridad (SSL, Cookies, HSTS, DEBUG)"),
        ("pip list --outdated", "Python Package Manager", "Detección de dependencias obsoletas en el entorno virtual"),
        ("npm audit (--prefix)", "Node Package Manager", "Detección de CVEs y vulnerabilidades en módulos JavaScript"),
        ("URL Resolver / Inspect", "php artisan route:list", "Mapeo exhaustivo de endpoints, permisos y decoradores"),
        ("git status & git ls-files", "Control de Versiones", "Verificación de fuga de credenciales o archivos sensibles versionados"),
        ("Postman v2.1 & Requests", "Pruebas de API REST", "Ejecución de batería de pruebas funcionales y de seguridad (HTTP Status)")
    ]
    for r in tool_rows:
        row = table_tools.add_row().cells
        row[0].text = r[0]
        row[1].text = r[1]
        row[2].text = r[2]

    # 5. MATRIZ DE AUDITORÍA DE DEPENDENCIAS
    doc.add_heading("5. Auditoría de Dependencias", level=1)
    doc.add_paragraph(
        "Se ejecutó una revisión sistemática de las dependencias instaladas tanto en el entorno Python (venv) "
        "como en los componentes Node.js existentes en el repositorio:"
    )

    table_dep = doc.add_table(rows=1, cols=5)
    table_dep.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_dep = table_dep.rows[0].cells
    hdr_dep[0].text = "Paquete"
    hdr_dep[1].text = "Versión Actual"
    hdr_dep[2].text = "Problema / Hallazgo"
    hdr_dep[3].text = "Severidad"
    hdr_dep[4].text = "Acción Recomendada"
    for c in hdr_dep:
        set_cell_background(c, "102C57")
        c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        c.paragraphs[0].runs[0].font.bold = True

    dep_data = [
        ("Django", "5.0.14", "Versión desactualizada frente a Django 6.x", "Media", "Actualizar a versión LTS con parches de seguridad"),
        ("djangorestframework", "3.16.0", "Versión previa a la 3.18.x", "Baja", "Actualizar manteniendo compatibilidad de ViewSets"),
        ("Pillow", "11.2.1", "Versión previa a 12.3.0", "Media", "Actualizar para parches de mitigación de corrupción de memoria"),
        ("tar (Node.js)", "<=7.5.20", "Arbitrary File Overwrite / Symlink Poisoning", "Crítica", "Ejecutar npm audit fix --force en banco-app"),
        ("axios (Node.js)", "1.0.0 - 1.19.0", "SSRF Bypass & Prototype Pollution Gadgets", "Alta", "Actualizar a la versión 1.20+"),
        ("path-to-regexp", "<0.1.13", "Regular Expression Denial of Service (ReDoS)", "Alta", "Actualizar dependencias de routing en Express")
    ]
    for r in dep_data:
        row = table_dep.add_row().cells
        row[0].text = r[0]
        row[1].text = r[1]
        row[2].text = r[2]
        row[3].text = r[3]
        row[4].text = r[4]

    # 6. REVISIÓN DE CONFIGURACIÓN Y SECRETOS
    doc.add_heading("6. Revisión de Configuración y Secretos", level=1)
    doc.add_paragraph(
        "Se analizaron los archivos de entorno, directivas de despliegue y control de versiones:\n"
        "1. Archivo .env y CREDENCIALES.txt: Correctamente incluidos en .gitignore, impidiendo que contraseñas reales se sincronicen en el repositorio remoto.\n"
        "2. SECRET_KEY: En settings.py existe un valor por defecto inseguro ('django-insecure-...'). En producción debe forzarse su provisión mediante variable de entorno estricta.\n"
        "3. Flags de Producción: DEBUG está activo en modo local, lo cual expone traceback interactivo con variables sensibles ante excepciones.\n"
        "4. Almacenamiento Git: Se detectaron archivos compilados (.pyc) y logs de depuración (logs/debug.log) previamente rastreados en git que deben ser eliminados del árbol con 'git rm --cached'."
    )

    # 7. MATRIZ GENERAL DE PRUEBAS TÉCNICAS (T01 - T21)
    doc.add_heading("7. Matriz General de Pruebas Técnicas", level=1)
    doc.add_paragraph("A continuación se presentan los resultados obtenidos tras la ejecución de las 21 pruebas automatizadas:")

    table_tests = doc.add_table(rows=1, cols=6)
    table_tests.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_t = table_tests.rows[0].cells
    hdr_t[0].text = "ID"
    hdr_t[1].text = "Prueba Ejecutada"
    hdr_t[2].text = "Resultado Esperado"
    hdr_t[3].text = "Resultado Obtenido"
    hdr_t[4].text = "HTTP"
    hdr_t[5].text = "Estado"
    for c in hdr_t:
        set_cell_background(c, "102C57")
        c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        c.paragraphs[0].runs[0].font.bold = True

    # Cargar datos de audit_results.json
    audit_file = "../logs/audit_results.json"
    if os.path.exists(audit_file):
        with open(audit_file, "r", encoding="utf-8") as f:
            t_data = json.load(f)
            for t in t_data:
                row = table_tests.add_row().cells
                row[0].text = t["id"]
                row[1].text = t["name"]
                row[2].text = t["expected"]
                row[3].text = t["status_desc"]
                row[4].text = str(t["http_code"])
                row[5].text = "APROBADA" if t["passed"] else "FALLIDA"
                fill_color = "D4EDDA" if t["passed"] else "F8D7DA"
                set_cell_background(row[5], fill_color)

    # 8. MATRIZ DE HALLAZGOS Y VULNERABILIDADES
    doc.add_heading("8. Matriz de Hallazgos y Vulnerabilidades", level=1)
    
    table_h = doc.add_table(rows=1, cols=4)
    table_h.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_h = table_h.rows[0].cells
    hdr_h[0].text = "ID"
    hdr_h[1].text = "Hallazgo Identificado"
    hdr_h[2].text = "Evidencia y Riesgo"
    hdr_h[3].text = "Recomendación / Mitigación"
    for c in hdr_h:
        set_cell_background(c, "102C57")
        c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        c.paragraphs[0].runs[0].font.bold = True

    hallazgos = [
        ("H01", "Acceso no autenticado a /dashboard/", "GET /dashboard/ devuelve HTTP 200 sin pedir login. Permite ver la interfaz gráfica a usuarios no registrados.", "Agregar LoginRequiredMixin a la vista en urls.py."),
        ("H02", "Exposición de listado de compras y ventas sin token", "FacturaCompraViewSet y VentaViewSet configuran AllowAny en list y retrieve. Información financiera expuesta.", "Configurar IsAuthenticated o permisos de rol específicos en list y retrieve."),
        ("H03", "Falta de control estricto de tipo y tamaño de archivo", "Subida de archivos falsos .exe responde 200 (error manejado silenciosamente). Archivos > 10MB son aceptados.", "Implementar validadores de extensión estricta, MIME type y tamaño máximo (2MB) con retorno HTTP 400/415."),
        ("H04", "Directivas de despliegue en modo inseguro", "python manage.py check --deploy reporta 6 advertencias (DEBUG=True, cookies sin flag Secure, HSTS ausente).", "Configurar SECURE_SSL_REDIRECT=True, SESSION_COOKIE_SECURE=True y CSRF_COOKIE_SECURE=True en producción."),
        ("H05", "Vulnerabilidades en dependencias NPM", "npm audit reportó 17 vulnerabilidades (1 crítica en tar, 12 altas en axios, lodash, etc.) en banco-app.", "Ejecutar npm audit fix y actualizar paquetes obsoletos en package.json."),
        ("H06", "Rastreo de artefactos y logs en Git", "Archivos compilados __pycache__ y debug.log se encuentran en git ls-files.", "Ejecutar 'git rm -r --cached __pycache__' y 'git rm --cached logs/debug.log'.")
    ]
    for h in hallazgos:
        row = table_h.add_row().cells
        row[0].text = h[0]
        row[1].text = h[1]
        row[2].text = h[2]
        row[3].text = h[3]

    # 9. GUION PARA VIDEO DE SOCIALIZACIÓN (5 - 8 MINUTOS)
    doc.add_heading("9. Guion Estructurado para el Video (5 - 8 Min)", level=1)
    doc.add_paragraph(
        "A continuación se presenta el guion técnico para la grabación del video explicativo conforme a los requerimientos de evaluación:\n\n"
        "• Minuto 0:00 - 0:45 | Presentación: Saludo, nombre del estudiante, carrera, materia Seguridad en Aplicaciones (MEDIT) y breve introducción de INBOT SYSTEM (API de inventarios con Django y SQL Server).\n"
        "• Minuto 0:45 - 1:45 | Comandos de Auditoría: Demostración en terminal ejecutando 'python manage.py check --deploy' y 'npm audit' mostrando los hallazgos automáticos.\n"
        "• Minuto 1:45 - 2:45 | Rutas y Mecanismos de Seguridad: Explicación de las rutas en urls.py, token authentication, CSRF middleware y RBAC en permissions.py.\n"
        "• Minuto 2:45 - 4:15 | Demostración en Postman (Pruebas Positivas y Negativas): Ejecución en vivo de Login correcto (200), Login erróneo (400), consulta de stock sin token (401) y prueba RBAC con Vendedor creando producto (403 Forbidden).\n"
        "• Minuto 4:15 - 5:30 | Prueba de Carga de Imágenes y Archivos: Demostración del endpoint de validación de imágenes y explicación del hallazgo de control de tamaño.\n"
        "• Minuto 5:30 - 6:30 | Hallazgo Principal y Mitigación: Explicar el hallazgo H01 o H02 y mostrar la corrección aplicada en el código.\n"
        "• Minuto 6:30 - 7:30 | Conclusiones y Aprendizaje: Reflexión sobre la importancia de la seguridad por diseño ('Security by Design') y cierre del video."
    )

    # 10. CONCLUSIONES
    doc.add_heading("10. Conclusiones", level=1)
    doc.add_paragraph(
        "La ejecución de esta auditoría integral sobre INBOT SYSTEM permitió corroborar que la aplicación posee "
        "una base sólida de seguridad basada en el ORM de Django (protección contra inyecciones SQL nativa), "
        "cifrado de contraseñas robusto (PBKDF2 SHA256), autenticación por tokens y protección CSRF efectiva en formularios web.\n\n"
        "No obstante, se identificaron áreas de mejora determinantes como la sobre-exposición de ciertos endpoints "
        "de lectura (list/retrieve en ventas y compras), la falta de restricción en vistas de plantillas web "
        "y la necesidad de endurecer las políticas de subida de imágenes y configuración de cookies seguras para entornos de producción. "
        "Este ejercicio demostró que una auditoría metódica es indispensable para transformar una aplicación funcional en una solución robusta y resistente a incidentes de seguridad."
    )

    doc.save("../scripts/Informe_Auditoria_Seguridad_Inbot_System.docx")
    print("Reporte DOCX generado exitosamente en scripts/Informe_Auditoria_Seguridad_Inbot_System.docx")

if __name__ == "__main__":
    create_report()
